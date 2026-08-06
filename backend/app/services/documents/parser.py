import os
import re
from abc import ABC, abstractmethod
from typing import List, Optional, Dict
import pdfplumber
from docx import Document as DocxDocument
import markdown
from bs4 import BeautifulSoup

from app.services.documents.metadata import ParsedDocument, ParsedPage
from app.services.documents.exceptions import ParsingError


class BaseParser(ABC):
    @abstractmethod
    def parse(self, file_path: str) -> ParsedDocument:
        pass


class PDFParser(BaseParser):
    def parse(self, file_path: str) -> ParsedDocument:
        try:
            pages = []
            headings = []
            all_text = ""

            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    all_text += text + "\n"

                    # Extract headings using font size heuristics
                    page_headings = []
                    for char in page.chars:
                        if char.get("size", 0) > 14 and char.get("text", "").strip():
                            page_headings.append({
                                "text": char["text"].strip(),
                                "page": i + 1,
                                "font_size": char["size"],
                                "bbox": char.get("bbox"),
                            })

                    pages.append(ParsedPage(
                        page_number=i + 1,
                        text=text,
                        headings=page_headings,
                    ))
                    headings.extend(page_headings)

            # Extract sections from headings
            sections = self._extract_sections(headings)

            return ParsedDocument(
                pages=pages,
                total_pages=len(pages),
                metadata={"parser": "pdfplumber"},
                headings=headings,
                sections=sections,
            )
        except Exception as e:
            raise ParsingError(f"PDF parsing failed: {str(e)}")

    def _extract_sections(self, headings: List[Dict]) -> List[Dict]:
        sections = []
        for h in headings:
            if h.get("font_size", 0) >= 14:
                sections.append({
                    "title": h["text"],
                    "page": h["page"],
                    "level": 1 if h["font_size"] >= 18 else 2,
                })
        return sections


class DocxParser(BaseParser):
    def parse(self, file_path: str) -> ParsedDocument:
        try:
            doc = DocxDocument(file_path)
            pages = []
            headings = []
            current_page = 1
            chars_per_page = 3000  # Approximate
            current_text = ""
            char_count = 0

            for para in doc.paragraphs:
                text = para.text
                if not text.strip():
                    continue

                style = para.style.name if para.style else ""

                # Detect headings
                if style.startswith("Heading"):
                    level = int(style.replace("Heading", "")) if style.replace("Heading", "").isdigit() else 1
                    headings.append({
                        "text": text,
                        "page": current_page,
                        "level": level,
                        "style": style,
                    })

                current_text += text + "\n"
                char_count += len(text)

                if char_count >= chars_per_page:
                    pages.append(ParsedPage(
                        page_number=current_page,
                        text=current_text,
                        headings=[h for h in headings if h.get("page") == current_page],
                    ))
                    current_page += 1
                    current_text = ""
                    char_count = 0

            # Add remaining text
            if current_text:
                pages.append(ParsedPage(
                    page_number=current_page,
                    text=current_text,
                    headings=[h for h in headings if h.get("page") == current_page],
                ))

            sections = self._extract_sections(headings)

            return ParsedDocument(
                pages=pages,
                total_pages=len(pages),
                metadata={"parser": "python-docx"},
                headings=headings,
                sections=sections,
            )
        except Exception as e:
            raise ParsingError(f"DOCX parsing failed: {str(e)}")

    def _extract_sections(self, headings: List[Dict]) -> List[Dict]:
        sections = []
        for h in headings:
            if h.get("level", 1) <= 3:
                sections.append({
                    "title": h["text"],
                    "page": h.get("page", 1),
                    "level": h.get("level", 1),
                })
        return sections


class TxtParser(BaseParser):
    def parse(self, file_path: str) -> ParsedDocument:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()

            # Split into pages based on form feeds or length
            pages_text = text.split("\f")
            if len(pages_text) == 1 and len(text) > 5000:
                # Split by approximate page size
                pages_text = [text[i:i+5000] for i in range(0, len(text), 5000)]

            pages = []
            for i, page_text in enumerate(pages_text):
                pages.append(ParsedPage(
                    page_number=i + 1,
                    text=page_text.strip(),
                ))

            return ParsedDocument(
                pages=pages,
                total_pages=len(pages),
                metadata={"parser": "txt"},
                headings=[],
                sections=[],
            )
        except Exception as e:
            raise ParsingError(f"TXT parsing failed: {str(e)}")


class MarkdownParser(BaseParser):
    def parse(self, file_path: str) -> ParsedDocument:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                md_text = f.read()

            # Convert to HTML for parsing
            html = markdown.markdown(md_text, extensions=['tables', 'fenced_code'])
            soup = BeautifulSoup(html, 'html.parser')

            # Extract headings
            headings = []
            for i, tag in enumerate(soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6'])):
                level = int(tag.name[1])
                headings.append({
                    "text": tag.get_text().strip(),
                    "page": 1,
                    "level": level,
                    "tag": tag.name,
                })

            # Split text into pages
            pages_text = md_text.split("\n\n")
            if len(pages_text) == 1:
                pages_text = [md_text]

            pages = []
            for i, page_text in enumerate(pages_text):
                pages.append(ParsedPage(
                    page_number=i + 1,
                    text=page_text.strip(),
                ))

            sections = self._extract_sections(headings)

            return ParsedDocument(
                pages=pages,
                total_pages=len(pages),
                metadata={"parser": "markdown"},
                headings=headings,
                sections=sections,
            )
        except Exception as e:
            raise ParsingError(f"Markdown parsing failed: {str(e)}")

    def _extract_sections(self, headings: List[Dict]) -> List[Dict]:
        sections = []
        for h in headings:
            if h.get("level", 1) <= 3:
                sections.append({
                    "title": h["text"],
                    "page": 1,
                    "level": h.get("level", 1),
                })
        return sections


def get_parser(file_type: str) -> BaseParser:
    parsers = {
        ".pdf": PDFParser(),
        ".docx": DocxParser(),
        ".txt": TxtParser(),
        ".md": MarkdownParser(),
    }
    parser = parsers.get(file_type.lower())
    if not parser:
        raise ParsingError(f"No parser available for file type: {file_type}")
    return parser